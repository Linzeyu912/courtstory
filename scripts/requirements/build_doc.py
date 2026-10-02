from pathlib import Path
import re
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE as RT

root=Path(__file__).resolve().parents[2]/'requirements/archive/broadcast-code-12'
src=(root/'赛事与报名资料.md').read_text(encoding='utf-8')
parts={}
for seg in src.split('\n## ')[1:]:
    name,body=seg.split('\n',1); parts[name]=body.strip()
doc=Document()
sec=doc.sections[0]; sec.page_width=Cm(21); sec.page_height=Cm(29.7)
sec.top_margin=Cm(1.8); sec.bottom_margin=Cm(1.8); sec.left_margin=Cm(2); sec.right_margin=Cm(2)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','List Bullet','List Number']:
    st=doc.styles[name]; st.font.name='Microsoft YaHei'; st.font.color.rgb=RGBColor(0,0,0)
    st.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Microsoft YaHei')
    st.font.size=Pt(10.5); st.paragraph_format.space_after=Pt(6); st.paragraph_format.line_spacing=1.15
doc.styles['Title'].font.size=Pt(22)
doc.styles['Heading 1'].font.size=Pt(15)
doc.styles['Heading 1'].paragraph_format.space_before=Pt(12)
doc.styles['Heading 2'].font.size=Pt(12)
footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('Broadcast Code 12  |  ')
f=OxmlElement('w:fldSimple'); f.set(qn('w:instr'),'PAGE'); footer._p.append(f)
for r in footer.runs:r.font.size=Pt(8)

def inline(p,s):
    pos=0
    for m in re.finditer(r'\[([^\]]+)\]\(([^)]+)\)',s):
        p.add_run(s[pos:m.start()])
        h=OxmlElement('w:hyperlink');h.set(qn('r:id'),p.part.relate_to(m[2],RT.HYPERLINK,is_external=True))
        r=OxmlElement('w:r'); pr=OxmlElement('w:rPr');c=OxmlElement('w:color');c.set(qn('w:val'),'175C91');pr.append(c);r.append(pr)
        t=OxmlElement('w:t');t.text=m[1];r.append(t);h.append(r);p._p.append(h);pos=m.end()
    p.add_run(s[pos:])

def table(lines):
    data=[[x.strip() for x in l.strip('|').split('|')] for l in lines if not re.match(r'^\|[- :|]+\|$',l)]
    t=doc.add_table(rows=0,cols=len(data[0]));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    widths=[2.9,14.1] if len(data[0])==2 else [1.0,9.0,7.0]
    for c,w in zip(t.columns,widths):c.width=Cm(w)
    borders=OxmlElement('w:tblBorders')
    for n in ['top','left','bottom','right','insideH','insideV']:
        b=OxmlElement('w:'+n);b.set(qn('w:val'),'single');b.set(qn('w:sz'),'4');b.set(qn('w:color'),'D9D9D9');borders.append(b)
    t._tbl.tblPr.append(borders)
    for i,row in enumerate(data):
        cells=t.add_row().cells
        trpr=t.rows[-1]._tr.get_or_add_trPr(); ns=OxmlElement('w:cantSplit');trpr.append(ns)
        if i==0:trpr.append(OxmlElement('w:tblHeader'))
        for j,(c,s) in enumerate(zip(cells,row)):
            c.width=Cm(widths[j]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=c._tc.get_or_add_tcPr();sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E5EDF4' if i==0 else ('F5F7F9' if i%2==0 else 'FFFFFF'));pr.append(sh)
            mar=OxmlElement('w:tcMar')
            for side,val in [('top','70'),('bottom','70'),('left','90'),('right','90')]:
                el=OxmlElement('w:'+side);el.set(qn('w:w'),val);el.set(qn('w:type'),'dxa');mar.append(el)
            pr.append(mar);p=c.paragraphs[0];p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.08
            if j==0:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            r=p.add_run(s);r.font.size=Pt(9.5);r.bold=(i==0)
    doc.add_paragraph().paragraph_format.space_after=Pt(0)

def content(text):
    lines=text.splitlines();i=0
    while i<len(lines):
        l=lines[i].strip()
        if l.startswith('|'):
            arr=[]
            while i<len(lines) and lines[i].strip().startswith('|'):arr.append(lines[i].strip());i+=1
            table(arr);continue
        if l:
            style='Normal'
            if l.startswith('- '):style='List Bullet';l=l[2:]
            elif re.match(r'^\d+\. ',l):style='List Number';l=re.sub(r'^\d+\. ','',l)
            inline(doc.add_paragraph(style=style),l)
        i+=1

def section(name):
    doc.add_heading(name,1);content(parts[name])
doc.add_paragraph('Broadcast Code 12\n黑客松赛事阅读手册',style='Title')
doc.add_paragraph('团队阅读与讨论资料',style='Subtitle')
doc.add_paragraph('信息采集日期  2026 年 9 月 17 日',style='Subtitle')
doc.add_paragraph('本次比赛要求用 AI 解读篮球比赛，并将数据分析、视觉标注和同步解说结合成可运行作品。报名截止为 9 月 23 日；取得参赛资格后，还需通过珠海初赛才能进入澳门 TOP 8 决赛。')
section('资格与规模');section('赛程')
doc.add_page_break()
section('比赛做什么');section('已公布的评审方向');section('作品交付')
doc.add_paragraph('阅读提示：报名筛选与作品评审是两个阶段。官网目前提供的是作品评审方向，尚不能据此推算报名入围概率。')
doc.add_page_break()
doc.add_heading('报名问题与材料要求',1)
field,terms=parts['报名表字段清单'].split('报名页显示需同意的条款：',1)
content(field)
doc.add_page_break()
section('奖项与保障')
doc.add_heading('报名条款',1);content(terms)
doc.add_page_break()
section('讨论前需要知道的未知项')
doc.add_heading('官方资料入口',1)
content('- [赛事官网及赛程](https://builderx.csdn.net/bcast-code?from=yaofoundation#eligibility)\n- [报名表](https://marketing.csdn.net/questions/Q2609071804379567473?from=yaofoundation)\n- [亚马逊云科技活动条款和条件](https://amazoncloud.cn/legal/eventterms)\n- [亚马逊云科技行为规范](https://amazoncloud.cn/legal/codeofconduct)')
doc.add_paragraph('报名表可从官网“立即报名”按钮进入。官网赛程以月日显示；本手册按当前 2026 年招募背景整理。赛程、奖品与活动安排以组委会最终通知为准。')
doc.core_properties.title='Broadcast Code 12 黑客松赛事阅读手册'
doc.core_properties.subject='赛事介绍 赛程 报名字段 评审方向'
doc.core_properties.author=''
out=root/'Broadcast Code 12 黑客松赛事阅读手册.docx'
doc.save(out)
print(out)
